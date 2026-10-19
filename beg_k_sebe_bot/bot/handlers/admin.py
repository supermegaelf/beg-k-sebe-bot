import io
import logging
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from aiogram import F, Router
from aiogram.filters import Command, CommandObject
from aiogram.types import BufferedInputFile, Message
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from beg_k_sebe_bot.bot.config import settings
from beg_k_sebe_bot.bot.database.models import Access, Payment, User
from beg_k_sebe_bot.bot.services.chat_access import remove_from_chats
from beg_k_sebe_bot.bot.services.payments import grant_access, notify_access_granted
from beg_k_sebe_bot.bot.services.tariffs import TARIFFS
from beg_k_sebe_bot.bot.texts import messages as msg
from beg_k_sebe_bot.bot.utils.program import today_msk

logger = logging.getLogger(__name__)
router = Router()
router.message.filter(F.from_user.id.in_(settings.admin_id_list))

_STATUS_LABELS = {"pending": "ожидает", "paid": "оплачен", "failed": "не прошёл"}
_KIND_LABELS = {"purchase": "покупка", "upgrade": "повышение"}
_SOURCE_LABELS = {"payment": "оплата", "manual": "вручную"}


def _dt(value: datetime | None) -> str:
    return value.astimezone(ZoneInfo(settings.timezone)).strftime("%Y-%m-%d %H:%M МСК") if value else ""


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in col) + 2)


@router.message(Command("payments"))
async def cmd_payments(message: Message, session: AsyncSession) -> None:
    users = {u.telegram_id: u for u in (await session.execute(select(User))).scalars().all()}
    payments = (await session.execute(select(Payment).order_by(Payment.inv_id))).scalars().all()
    accesses = (await session.execute(select(Access))).scalars().all()
    accesses = sorted(accesses, key=lambda a: (-TARIFFS[a.tariff].rank, a.granted_at))

    def username(user_id: int) -> str:
        user = users.get(user_id)
        return f"@{user.username}" if user and user.username else ""

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Доступы", ["Telegram ID", "Юзернейм", "Тариф", "Доступ до", "Источник", "Выдан"], [
        [str(a.user_id), username(a.user_id), TARIFFS[a.tariff].title, a.access_until.isoformat(),
         _SOURCE_LABELS.get(a.source, a.source), _dt(a.granted_at)]
        for a in accesses
    ])
    _sheet(wb, "Платежи", ["Счёт", "Telegram ID", "Юзернейм", "Тариф", "Тип", "Сумма, ₽", "Промокод", "Статус", "Создан", "Оплачен"], [
        [p.inv_id, str(p.user_id), username(p.user_id), TARIFFS[p.tariff].title, _KIND_LABELS.get(p.kind, p.kind),
         p.amount, p.promo_code or "", _STATUS_LABELS.get(p.status, p.status), _dt(p.created_at), _dt(p.paid_at)]
        for p in payments
    ])

    bio = io.BytesIO()
    wb.save(bio)
    await message.answer_document(
        BufferedInputFile(bio.getvalue(), filename=f"payments_{today_msk().isoformat()}.xlsx"),
        caption=f"Доступов: {len(accesses)} · платежей: {len(payments)}",
    )


@router.message(Command("grant"))
async def cmd_grant(message: Message, command: CommandObject, session: AsyncSession) -> None:
    parts = (command.args or "").split()
    if len(parts) != 2 or not parts[0].isdigit() or parts[1] not in TARIFFS:
        await message.answer(msg.ADMIN_GRANT_USAGE)
        return
    user_id, tariff = int(parts[0]), parts[1]

    user = await session.get(User, user_id)
    if user is None:
        user = User(telegram_id=user_id, joined_at=datetime.now(timezone.utc))
        session.add(user)
        await session.commit()

    access = grant_access(session, await session.get(Access, user_id), user_id, tariff, "manual")
    await session.commit()
    await notify_access_granted(message.bot, user, access)
    await message.answer(msg.ADMIN_GRANT_OK.format(user_id=user_id, tariff=TARIFFS[tariff].title))


@router.message(Command("revoke"))
async def cmd_revoke(message: Message, command: CommandObject, session: AsyncSession) -> None:
    arg = (command.args or "").strip()
    if not arg.isdigit():
        await message.answer(msg.ADMIN_REVOKE_USAGE)
        return
    user_id = int(arg)

    access = await session.get(Access, user_id)
    if access is None:
        await message.answer(msg.ADMIN_REVOKE_NONE.format(user_id=user_id))
        return
    await session.delete(access)
    await session.commit()

    await remove_from_chats(message.bot, user_id)
    try:
        await message.bot.send_message(user_id, msg.ACCESS_REVOKED)
    except Exception as e:
        logger.warning("Failed to notify %d about revoked access: %s", user_id, e)
    await message.answer(msg.ADMIN_REVOKE_OK.format(user_id=user_id))
