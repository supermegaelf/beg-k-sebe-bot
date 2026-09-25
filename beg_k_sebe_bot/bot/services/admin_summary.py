import io
import logging
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from aiogram import Bot
from aiogram.types import BufferedInputFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from beg_k_sebe_bot.bot.config import settings
from beg_k_sebe_bot.bot.database.models import DailyCheckin, SentEvent, User, WeeklyReflection
from beg_k_sebe_bot.bot.utils.program import today_msk

logger = logging.getLogger(__name__)

_PROGRAM_DAYS = settings.final_program_day - 1

# Plain format names for the export (no emoji, one word).
_FORMAT_NAME = {"walk_22min": "ходьба", "run_22min": "бег", "own": "другое"}

_WEEK_LEN = 7

# (заголовок, ширина колонки). Порядок колонок совпадает с записью в _build_xlsx().
# Раскладка: блок на участника (разделён рамкой), внутри — строки-дни. Поля уровня
# участника и уровня недели объединяются по вертикали (merge), поэтому не дублируются
# и ничего не перезаписывается — история копится строками-днями.
_COLUMNS = [
    ("Telegram ID", 14),                        # 1  участник
    ("Юзернейм", 16),                           # 2  участник
    ("Формат движения", 13),                    # 3  участник
    ("Точка А: оценка (1–10)", 10),             # 4  участник
    ("Точка А: описание", 30),                  # 5  участник
    ("Точка Б: оценка (1–10)", 10),             # 6  участник
    ("Точка Б: описание", 30),                  # 7  участник
    ("Неделя", 7),                              # 8  неделя
    ("Дни", 6),                                 # 9  день
    ("Чек-инов сделано", 10),                   # 10 неделя
    ("Чек-инов пропущено", 10),                 # 11 неделя
    ("% выполнения", 10),                       # 12 неделя
    ("Ходьба, мин", 10),                        # 13 день
    ("Бег, мин", 9),                            # 14 день
    ("Другое, мин", 10),                        # 15 день
    ("Всего минут движения", 12),               # 16 день
    ("Ежедневно: что помогло", 40),             # 17 день
    ("Ежедневно: что было сложнее всего", 40),  # 18 день
    ("Ежедневно: что сдвинулось к цели", 40),   # 19 день
    ("Рефлексия: результат недели", 40),        # 20 неделя
    ("Рефлексия: что помогало", 34),            # 21 неделя
    ("Рефлексия: сложнее всего", 34),           # 22 неделя
    ("Рефлексия: ближе к цели", 34),            # 23 неделя
    ("Рефлексия: чем поделиться", 34),          # 24 неделя
]

# Колонки, объединяемые по вертикали на весь блок участника / на неделю.
_PERSON_COLS = (1, 2, 3, 4, 5, 6, 7)
_WEEK_COLS = (8, 10, 11, 12, 20, 21, 22, 23, 24)


def _week_bounds(week):
    first_day = (week - 1) * _WEEK_LEN + 1
    last_day = min(week * _WEEK_LEN, _PROGRAM_DAYS)
    return first_day, last_day


def _week_agg(checkins, first_day, last_day, user_start, today):
    """Аггрегаты недели: сделано / пропущено / % (минуты теперь по дням, не тут)."""
    prog_end = settings.start_date + timedelta(days=_PROGRAM_DAYS - 1)
    week_dates = [settings.start_date + timedelta(days=d - 1) for d in range(first_day, last_day + 1)]
    answered = [c for c in checkins
                if c.status == "answered" and first_day <= c.day_number <= last_day]

    # Пропущенными считаем только полностью прошедшие дни (сегодня ещё можно сделать).
    last_full = min(today - timedelta(days=1), prog_end)
    expected_past = sum(1 for d in week_dates if user_start <= d <= last_full)
    completed_past = sum(1 for c in answered if c.date < today)
    missed = max(0, expected_past - completed_past)

    last_incl = min(today, prog_end)
    expected_incl = sum(1 for d in week_dates if user_start <= d <= last_incl)
    completed = len(answered)
    pct = min(round(completed / expected_incl * 100), 100) if expected_incl > 0 else 0
    return completed, missed, pct


def _person_weeks(user, checkins, reflections, today) -> list[dict]:
    """Структура блока участника: список недель, в каждой — строки-дни (день 1..7)."""
    onboarding_date = user.onboarding_completed_at.date() if user.onboarding_completed_at else settings.start_date
    onboarding_day = max(1, (onboarding_date - settings.start_date).days + 1)
    user_start = max(settings.start_date, onboarding_date)
    cur_day = min(max((today - settings.start_date).days + 1, 1), _PROGRAM_DAYS)
    start_week = (onboarding_day - 1) // _WEEK_LEN + 1
    current_week = (cur_day - 1) // _WEEK_LEN + 1
    checkin_by_day = {c.day_number: c for c in checkins if c.status == "answered"}
    refl_by_week = {r.week_number: r for r in reflections}

    weeks = []
    for week in range(start_week, current_week + 1):
        first_day, last_day = _week_bounds(week)
        days = [d for d in range(first_day, last_day + 1) if onboarding_day <= d <= cur_day]
        if not days:
            continue
        done, missed, pct = _week_agg(checkins, first_day, last_day, user_start, today)
        day_rows = []
        for d in days:
            c = checkin_by_day.get(d)
            mins = c.minutes if c and c.minutes else 0
            cat = c.activity_category if c and c.minutes else None
            day_rows.append({
                "day": d,
                "walk": mins if cat == "walk" else 0,
                "run": mins if cat == "run" else 0,
                "own": mins if cat == "own" else 0,
                "total": mins,
                "help": (c.help_text if c else "") or "",
                "hardest": (c.hardest_text if c else "") or "",
                "shift": (c.shift_text if c else "") or "",
            })
        weeks.append({
            "week": week, "done": done, "missed": missed, "pct": pct,
            "reflection": refl_by_week.get(week), "days": day_rows,
        })
    return weeks


def _build_xlsx(users, checkins_by_user, reflections_by_user, today) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Участники по неделям"

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    person_sep = Side(style="medium", color="4472C4")  # рамка-разделитель между участниками
    cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for col, (title, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.column_dimensions[get_column_letter(col)].width = width

    n_cols = len(_COLUMNS)
    merges = []  # (r1, c1, r2, c2) — применяем после записи всех ячеек
    row = 2
    for user in users:
        weeks = _person_weeks(user, checkins_by_user.get(user.telegram_id, []),
                              reflections_by_user.get(user.telegram_id, []), today)
        if not weeks:
            continue
        person_first = row
        person_vals = {
            1: str(user.telegram_id),
            2: user.username or "",
            3: _FORMAT_NAME.get(user.movement_format, user.movement_format or ""),
            4: user.point_a_score if user.point_a_score is not None else "",
            5: user.point_a_text or "",
            6: user.point_b_score if user.point_b_score is not None else "",
            7: user.point_b_text or "",
        }
        for wk in weeks:
            week_first = row
            r = wk["reflection"]
            week_vals = {
                8: wk["week"], 10: wk["done"], 11: wk["missed"], 12: wk["pct"],
                20: (r.result_text if r else "") or "",
                21: (r.helped_text if r else "") or "",
                22: (r.hardest_text if r else "") or "",
                23: (r.progress_text if r else "") or "",
                24: (r.share_text if r else "") or "",
            }
            for i, dr in enumerate(wk["days"]):
                values = {
                    9: dr["day"], 13: dr["walk"], 14: dr["run"], 15: dr["own"], 16: dr["total"],
                    17: dr["help"], 18: dr["hardest"], 19: dr["shift"],
                }
                if i == 0:
                    values.update(week_vals)
                if row == person_first:
                    values.update(person_vals)
                top = person_sep if row == person_first else thin
                for col in range(1, n_cols + 1):
                    cell = ws.cell(row=row, column=col, value=values.get(col, ""))
                    cell.alignment = cell_align
                    cell.border = Border(left=thin, right=thin, top=top, bottom=thin)
                row += 1
            if row - 1 > week_first:  # объединяем поля недели по её дням
                merges += [(week_first, c, row - 1, c) for c in _WEEK_COLS]
        if row - 1 > person_first:  # объединяем поля участника по всему блоку
            merges += [(person_first, c, row - 1, c) for c in _PERSON_COLS]

    for r1, c1, r2, c2 in merges:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "A2"

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


async def send_admin_summary(bot: Bot, session: AsyncSession) -> None:
    admin_ids = settings.admin_id_list
    if not admin_ids:
        logger.warning("ADMIN_IDS not set, skipping admin summary")
        return

    today = today_msk()
    marker_key = f"admin_summary:{today.isoformat()}"
    if await session.get(SentEvent, marker_key) is not None:
        logger.info("Admin summary already sent for %s, skipping", today)
        return

    users = (await session.execute(
        select(User).where(User.onboarding_completed_at.is_not(None))
    )).scalars().all()
    if not users:
        return
    user_ids = [u.telegram_id for u in users]

    checkins = (await session.execute(
        select(DailyCheckin).where(DailyCheckin.user_id.in_(user_ids))
    )).scalars().all()
    checkins_by_user: dict[int, list] = defaultdict(list)
    for c in checkins:
        checkins_by_user[c.user_id].append(c)

    reflections = (await session.execute(
        select(WeeklyReflection).where(
            WeeklyReflection.user_id.in_(user_ids),
            WeeklyReflection.status == "answered",
        )
    )).scalars().all()
    reflections_by_user: dict[int, list] = defaultdict(list)
    for r in reflections:  # keep every answered week so history accumulates
        reflections_by_user[r.user_id].append(r)

    data = _build_xlsx(users, checkins_by_user, reflections_by_user, today)
    filename = f"beg_k_sebe_{today.isoformat()}.xlsx"

    sent = False
    for admin_id in admin_ids:
        try:
            await bot.send_document(
                admin_id,
                BufferedInputFile(data, filename=filename),
                caption=f"Сводка по участникам на {today.isoformat()} ({len(users)} чел.)",
            )
            sent = True
        except Exception as e:
            logger.error("Failed to send admin summary to %d: %s", admin_id, e)

    if sent:
        session.add(SentEvent(key=marker_key, sent_at=datetime.now(timezone.utc)))
        await session.commit()
    logger.info("Admin summary sent to %d admins for %d users", len(admin_ids), len(users))
